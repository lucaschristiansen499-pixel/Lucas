from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC_PATH = "Lucas_Notecard_QuickGuide.docx"
LINK = "https://claude.ai/artifact/W9gNFyg4QHN7y6i1Mq3QUf"

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

for s in doc.sections:
    s.left_margin = Inches(0.75)
    s.right_margin = Inches(0.75)
    s.top_margin = Inches(0.6)
    s.bottom_margin = Inches(0.6)


def set_spacing(p, before=0, after=6):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)


def title(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1C, 0x24, 0x31)
    set_spacing(p, 0, 2)
    return p


def subtitle(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x5B, 0x64, 0x72)
    set_spacing(p, 0, 10)
    return p


def section_head(number, text):
    p = doc.add_paragraph()
    r = p.add_run(f"{number}.  {text}")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x2A, 0x5D, 0xB0)
    set_spacing(p, 10, 4)
    return p


def step(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    set_spacing(p, 0, 4)
    p.paragraph_format.left_indent = Inches(0.25)
    return p


def note(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x5B, 0x64, 0x72)
    set_spacing(p, 2, 8)
    return p


def hr():
    p = doc.add_paragraph()
    set_spacing(p, 4, 4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)


# ============================================================
title("Lucas's Notecard — Quick Guide")
subtitle("Two tools, one page. It runs on your own Claude account — nothing is sent anywhere else.")

p = doc.add_paragraph()
r = p.add_run("Open it here: ")
r.bold = True
r2 = p.add_run(LINK)
r2.font.color.rgb = RGBColor(0x2A, 0x5D, 0xB0)
set_spacing(p, 0, 8)

hr()

# ------------------------------------------------------------
section_head(1, "Turn lecture notes into a study card")
step("Click the “Lecture notes” tab (it's already selected when the page opens).")
step("Paste your AI Pocket transcript or lecture notes into the box.")
step("Click “Organize into a study card.”")
step("Wait a few seconds — the card fills in with a plain-language summary, key terms, steps, examples, and a homework checklist.")
step("Click “Save card (.md)” if you want to keep a copy on your device.")
note("If part of the source was unclear, the card will say so under “Double-check this” instead of guessing — go check that part yourself.")

# ------------------------------------------------------------
section_head(2, "Break a task into steps")
step("Click the “Break it down” tab at the top.")
step("Paste in an assignment, or just type out whatever's stuck in your head about it.")
step("Click “Break into steps.”")
step("You'll get a short checklist — check off each step as you finish it.")
step("Click “Save list (.txt)” to keep a copy.")

hr()

p = doc.add_paragraph()
r = p.add_run("Good to know: ")
r.bold = True
p.add_run("Your last input and output are saved automatically in the browser, so reopening the page brings back your last card or list. If a button looks greyed out, Claude access hasn't been turned on for the page yet — ask whoever shared the link with you to check the Share settings.")
set_spacing(p, 4, 0)

doc.save(DOC_PATH)
print("Saved:", DOC_PATH)
