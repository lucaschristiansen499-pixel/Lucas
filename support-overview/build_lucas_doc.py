from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC_PATH = "Lucas_Support_Overview.docx"

doc = Document()

# ---- base style ----
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

sections = doc.sections
for s in sections:
    s.left_margin = Inches(1)
    s.right_margin = Inches(1)
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)


def add_title(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x1C, 0x24, 0x31)
    p.space_after = Pt(2)
    return p


def add_subtitle(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x5B, 0x64, 0x72)
    return p


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def body(text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def sub_bullet(text):
    p = doc.add_paragraph(style="List Bullet 2")
    p.add_run(text)
    return p


def numbered(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()


# ============================================================
add_title("Lucas's Support Overview")
add_subtitle("A working reference for medical, school, and home support — compiled September 2026")
hr()

# ------------------------------------------------------------
h1("1. Background")
body(
    "Lucas Christiansen (DOB 11/28/2007) was diagnosed in October 2023 with metastatic "
    "(leptomeningeal) Group 3 medulloblastoma (WHO Grade IV). Treatment at Phoenix Children's "
    "Hospital included surgical resection (10/16/23), chemoradiotherapy (completed 12/2023), and "
    "six cycles of maintenance chemotherapy (completed 06/2024). The family relocated from Arizona "
    "to Illinois in August 2025; Lucas is now a junior (Class of 2028) at Marmion Academy in Aurora, "
    "IL, and his ongoing care has transferred to Lurie Children's Hospital in Chicago."
)

h2("Ongoing effects being managed")
bullet("Peripheral neuropathy in feet/legs and balance issues — continues physical therapy twice weekly.")
bullet("Right eye was patched for several months post-surgery for double vision (now largely resolved); uses blue light glasses.")
bullet("Thyroid damage — daily thyroid medication.")
bullet("Hormonal effects — weekly testosterone injections (since April 2025) and daily growth hormone injections (since August 2025).")
bullet("Mild right-sided, high-frequency hearing loss — uses custom protective hearing aids.")
bullet("Documented mood lability and irritability tied to frustration with his recovery pace, particularly around golf performance and schoolwork.")

# ------------------------------------------------------------
h1("2. Neuropsychological Evaluation (Baseline — 12/30/2024)")
body(
    "Conducted at Phoenix Children's Hospital by Dr. Laura Winstone-Weide, Pediatric Neuropsychologist. "
    "Diagnosis: Mild Neurocognitive Disorder due to another medical condition (medulloblastoma and its "
    "treatment). A second evaluation has recently been completed; results are pending as of this writing."
)

h2("Key results")
add_table(
    ["Domain", "Result"],
    [
        ["Full Scale IQ", "77 (Below Average, 6th percentile)"],
        ["Verbal Comprehension", "87 — relative strength"],
        ["Fluid Reasoning", "72 — weakness"],
        ["Working Memory", "78 — weakness"],
        ["Processing Speed", "84 — age-appropriate, but flagged for monitoring (late effects can emerge years later)"],
        ["Fine motor speed/dexterity", "<1st percentile, both hands"],
        ["Visual-motor integration", "<1st percentile"],
        ["Verbal learning & memory", "75th-84th percentile — a clear strength"],
        ["Spelling", "3rd percentile"],
    ],
)

h2("What this means day to day")
bullet("Strong verbal reasoning and verbal memory — Lucas learns best through spoken, plain-language explanation.")
bullet("Significant difficulty with executive function: working memory, planning, and shifting between tasks.")
bullet("Marked difficulty translating visual information into precise hand movements — affects handwriting, note-taking speed, and (by his own report) golf swing accuracy.")
bullet("Homework, especially math, reportedly takes about twice as long as it does for classmates.")

h2("Original recommendations most relevant to daily support")
bullet("Break multi-step problems into explicit steps; provide notecards listing steps for a given problem type.")
bullet("Allow recording of lectures and photographing the whiteboard, since real-time note-taking is a documented weak point.")
bullet("Provide a laptop or dictation option for written work given the fine-motor/handwriting difficulty.")
bullet("Preferential seating, extended time (time and a half), and a quiet testing space.")
bullet("An executive-function study hall or equivalent organizational support.")
bullet("Continued PT/OT, counseling support for frustration and anxiety, and a neuropsychological re-evaluation at roughly the two-year mark (now underway).")

# ------------------------------------------------------------
h1("3. Current School Accommodations — Marmion Academy")
body(
    "Marmion's Student Support Plan (created 10/9/2025; school counselor Christina Lipp) currently includes:"
)
bullet("Preferential seating near instruction / away from distractions.")
bullet("Teacher-provided notes to supplement (not replace) Lucas's own notes.")
bullet("Extended time on tests, up to time and a half; option to test in the Academic Center.")
bullet("Breaks as needed on longer assessments.")
bullet("Weekly check-ins with his school counselor.")

h2("Gaps compared to the original neuropsychological recommendations")
bullet("No math-specific scaffolding — no formula sheets or step-by-step notecards for problem types.")
bullet("No handwriting/dictation accommodation (laptop use, dictation software) despite documented fine-motor and visual-motor deficits.")
bullet("No executive-function study hall or organizational support built in.")
bullet("The \"Physical/Medical Support\" section of the SSP is currently blank.")
body(
    "These gaps are worth revisiting once Lurie Children's provides updated recommendations from the "
    "second neuropsychological evaluation.",
    italic=True,
)

# ------------------------------------------------------------
h1("4. In Progress: AI Pocket Note-Taking Accommodation")
body(
    "The family has provided Lucas with an AI Pocket note-taking device to address his documented "
    "note-taking difficulty, directly matching the original evaluation's recommendation to let him "
    "record lectures. Two requests are pending with Marmion:"
)
numbered("Formally add use of the device to Lucas's Student Support Plan under Instructional Support.")
numbered("Grant a medical-necessity exception to Marmion's no-cell-phone-during-school-hours policy specifically for this device.")

h2("Draft request letter (to Christina Lipp, School Counselor)")
body(
    "We wanted to follow up on Lucas's Student Support Plan (dated 10/9/2025) to request an additional "
    "accommodation.",
)
body(
    "Lucas's neuropsychological evaluation (Phoenix Children's Hospital, 12/30/2024) documented "
    "significant weaknesses in fine motor speed/dexterity and visual-motor integration (both below "
    "the 1st percentile), along with working memory difficulties. The evaluating neuropsychologist "
    "specifically recommended that Lucas be allowed to record lectures and photograph the whiteboard "
    "so he can revisit material afterward, noting he “may have significant difficulty with "
    "efficient/accurate note-taking during lectures or class discussions.”"
)
body(
    "To support this documented need, we've provided Lucas with an AI Pocket note-taking device, used "
    "specifically to record and transcribe classroom instruction so he can review it later. We'd like "
    "to request:"
)
numbered("That use of this device during instructional time be formally added to Lucas's Student Support Plan under Instructional Support, consistent with his neuropsychologist's recommendation.")
numbered("A medical-necessity exception to Marmion's no-cell-phone-during-school-hours policy specifically for this device.")
body(
    "We're currently waiting on updated recommendations from Lucas's care team at Lurie Children's "
    "Hospital and will forward that documentation as soon as it's available to further support this "
    "request. In the meantime, we're happy to provide the original neuropsychological report, or the "
    "relevant excerpt, if that would help formalize the accommodation."
)
body("Thank you for your continued partnership in supporting Lucas.")
body("Best,\nJon Christiansen")
body(
    "Note: a brief supporting note from Lucas's Lurie's doctor, once available, would strengthen the "
    "phone-policy exception request significantly.",
    italic=True,
)

# ------------------------------------------------------------
h1("5. Reminders: PlusPortals → Google Calendar")
body(
    "Marmion uses PlusPortals (Rediker Software) as its student information system, which has a "
    "built-in Google Calendar Feed export. Setting this up gives Lucas automatic reminders for "
    "assignment due dates without adding any new app or data-privacy exposure."
)
h2("On the PlusPortals side")
numbered("Log into PlusPortals (parent or student account).")
numbered("Go to the Calendar tab/module.")
numbered("Look for a “Subscribe,” “Feed,” or “Sync to Google” icon/link near the top of the calendar view, and copy the generated feed URL.")
body("If it isn't visible, ask Marmion's registrar/front office to enable the “Google Calendar Feed” setting.", italic=True)

h2("On the Google Calendar side (from a browser, not the phone app)")
numbered("Go to calendar.google.com, signed into the account Lucas uses on his phone.")
numbered("Click the “+” next to “Other calendars” in the left sidebar.")
numbered("Choose “From URL” and paste the PlusPortals feed link (swap webcal:// for https:// if needed).")
numbered("Click “Add calendar,” then set default notifications on it, and confirm it's visible in the Google Calendar app on Lucas's phone.")
body(
    "Note: Google refreshes subscribed calendars every several hours, not instantly. Also turn on push "
    "notifications in the free PlusPortals mobile app as a faster backup.",
    italic=True,
)

# ------------------------------------------------------------
h1("6. Digital Toolkit — “Lucas's Notecard”")
body(
    "To keep everything inside a tool we control (rather than a third-party app with broad account "
    "access), a private tool was built directly in Claude:"
)
p = doc.add_paragraph()
p.add_run("Link: ").bold = True
p.add_run("https://claude.ai/artifact/W9gNFyg4QHN7y6i1Mq3QUf")
bullet("Lecture Notes mode: paste an AI Pocket transcript or summary; it reorganizes it into a study card (plain-language summary, simplified terms, numbered steps, worked examples, a flagged “unclear — don't guess” section, and a homework checklist).")
bullet("Break It Down mode: paste an assignment or a brain-dump; it returns a short, checkable list of concrete steps.")
bullet("Both modes can save their output as a file. It runs on Lucas's/the viewer's own Claude account — no third-party company, no new data-sharing relationship.")
body(
    "The page is private by default and needs to be shared with Lucas's account before he can use it "
    "independently (via the page's Share menu).",
    italic=True,
)

# ------------------------------------------------------------
h1("7. Evaluated and Not Recommended: General AI Assistant Apps")
body(
    "An app called Instinct (app.instinct.com) was evaluated as a possible executive-function support "
    "tool and set aside for now. It is a broad, autonomous AI agent that connects to email, messaging, "
    "screen, audio, and location, and can independently take actions (including purchases) on a user's "
    "behalf. Concerns: very broad account access for a young adult with documented executive-function "
    "and judgment challenges; it trains its AI on user data by default unless manually opted out; "
    "liability is capped at $100; and it is a brand-new, unproven company. Lower-risk, purpose-built "
    "alternatives (the calendar sync and the Notecard tool above) address the same needs without that risk."
)

# ------------------------------------------------------------
h1("8. Open Items")
bullet("Second neuropsychological evaluation results — pending; doctors at Lurie Children's have shared some preliminary feedback (memory issues, further decline in age-expected executive function).")
bullet("Waiting on Lurie Children's formal recommendations, which may support updates to Lucas's Student Support Plan and the phone-policy exception request.")
bullet("Marmion SSP update — revisit the gaps noted in Section 3 once Lurie's recommendations arrive.")
bullet("AI Pocket / phone-policy request — send the letter in Section 4; a supporting note from Lurie's would help.")
bullet("Decide whether the “Break It Down” task list in the Notecard tool should be shared between Lucas and a parent, or stay private to Lucas.")

doc.save(DOC_PATH)
print("Saved:", DOC_PATH)
